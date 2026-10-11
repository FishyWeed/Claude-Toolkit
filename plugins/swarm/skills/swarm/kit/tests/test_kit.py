"""swarm_kit smoke test: builds a chart, a deck, a document, a video and the gallery in a temp folder.

Run with the swarm environment's Python:  python tests/test_kit.py [--keep]
"""

from __future__ import annotations

import json
import shutil
import sys
import tempfile
from pathlib import Path


def main() -> int:
    keep = "--keep" in sys.argv
    root = Path(tempfile.mkdtemp(prefix="swarm-kit-test-"))
    try:
        import matplotlib

        matplotlib.use("Agg")
        import matplotlib.pyplot as plt

        from swarm_kit import deck, doc, gallery, video
        from swarm_kit.themes import THEMES, apply_matplotlib

        (root / "charts").mkdir()
        apply_matplotlib("ocean")
        fig, ax = plt.subplots()
        ax.bar(["2022", "2023", "2024", "2025"], [1.2, 1.9, 2.6, 3.4])
        ax.set_title("Revenue grew almost 3x in three years")
        ax.set_ylabel("Revenue ($M)")
        fig.savefig(root / "charts" / "revenue.png")
        plt.close(fig)

        from PIL import Image, ImageDraw

        (root / "images").mkdir()
        bg = Image.new("RGB", (2400, 1350), (40, 90, 60))
        d = ImageDraw.Draw(bg)
        for y in range(1350):
            d.line([(0, y), (2400, y)], fill=(40 + y // 12, 90 + y // 20, 60 + y // 15))
        d.ellipse([1500, 300, 2100, 900], fill=(240, 200, 120))
        bg.save(root / "images" / "bg.png")

        (root / "slides").mkdir()
        spec = {
            "title": "Smoke Test Deck", "subtitle": "Every layout once", "author": "swarm_kit", "date": "2026-10-10", "theme": "ocean",
            "sources": ["[1] Example source — Publisher, 2025. https://example.com"],
            "slides": [
                {"layout": "title", "image": "../images/bg.png", "kicker": "Background image"},
                {"layout": "section", "kicker": "Part 1", "title": "The opportunity", "image": "../images/bg.png"},
                {"layout": "section", "kicker": "Part 2", "title": "A plain section"},
                {"layout": "bullets", "title": "Three reasons this works", "bullets": ["**Fast**: ships in a week", "*Cheap* to run", "Uses `ffmpeg` and [Edge](https://example.com)"], "notes": "Say why.", "source": "[1]"},
                {"layout": "chart", "title": "Revenue grew almost 3x", "image": "../charts/revenue.png", "caption": "Revenue, 2022-2025"},
                {"layout": "image-text", "title": "A picture and words", "image": "../charts/revenue.png", "bullets": ["One", "Two"]},
                {"layout": "two-column", "title": "Before and after", "left": {"heading": "Before", "bullets": ["Slow", "Manual"]}, "right": {"heading": "After", "bullets": ["Fast", "Automatic"]}},
                {"layout": "stats", "title": "By the numbers", "stats": [{"value": "3.4M", "label": "Revenue 2025", "source": "[1]"}, {"value": "42%", "label": "Growth"}, {"value": "12", "label": "Markets"}]},
                {"layout": "quote", "quote": "The best way to predict the future is to build it.", "attribution": "Someone"},
                {"layout": "table", "title": "Options compared", "columns": ["Option", "Cost", "Speed"], "rows": [["A", "$", "Fast"], ["B", "$$", "Faster"]]},
                {"layout": "closing", "title": "The ask", "bullets": ["Approve the pilot", "Two engineers for a quarter"]},
                {"layout": "closing", "title": "Closing with a picture", "bullets": ["One", "Two"], "image": "../images/bg.png"},
                {"layout": "sources"},
            ],
        }
        (root / "slides" / "deck.json").write_text(json.dumps(spec), encoding="utf-8")
        result = deck.build(root / "slides" / "deck.json")
        assert Path(result["pptx"]).stat().st_size > 20_000, result
        page = Path(result["html"]).read_text(encoding="utf-8")
        assert page.count("has-bg") >= 3, "title and section backgrounds missing"
        assert 'layout-image-text' in page, "closing with an image should lay out as image-text"
        from pptx import Presentation

        prs = Presentation(result["pptx"])
        first = prs.slides[0]
        assert any(sh.shape_type == 13 for sh in first.shapes), "title background picture missing in PPTX"
        assert Path(result["html"]).stat().st_size > 10_000, result
        assert not result["warnings"] or all("PDF" in w for w in result["warnings"]), result["warnings"]
        print("deck ok:", {k: v for k, v in result.items() if k != "warnings"}, result["warnings"])

        (root / "docs").mkdir()
        (root / "docs" / "summary.md").write_text(
            "---\ntitle: Executive Summary\nsubtitle: Smoke test\nauthor: swarm_kit\ndate: 2026-10-10\ntheme: midnight\n---\n\n"
            "## The short version\n\nRevenue grew **3x** in three years [1]. See [the source](https://example.com).\n\n"
            "1. First numbered\n2. Second numbered\n   - nested bullet\n\n- A bullet\n- Another\n\n"
            "| Metric | 2024 | 2025 |\n|---|---|---|\n| Revenue | 2.6 | 3.4 |\n| Margin | 18% | 22% |\n\n"
            "![Revenue by year](../charts/revenue.png)\n\n> A quote worth reading.\n\n```\ncode block\n```\n\n"
            "1. A new list restarts at one\n\n---\n\n## Sources\n\n[1] Example — Publisher, 2025.\n",
            encoding="utf-8",
        )
        result = doc.build(root / "docs" / "summary.md", formats=("docx", "pdf", "html"))
        assert Path(result["docx"]).stat().st_size > 20_000, result
        assert result["theme"] == "ocean", "dark themes must print on white"
        print("doc ok:", {k: v for k, v in result.items() if k != "warnings"}, result["warnings"])

        (root / "video").mkdir()
        vspec = {
            "size": "1280x720", "fps": 24, "theme": "midnight", "transition": 0.5,
            "title": {"text": "Smoke test", "subtitle": "swarm_kit video", "duration": 2},
            "slides": [{"image": "../charts/revenue.png", "duration": 2.5, "caption": "Revenue grew almost 3x", "motion": "zoom-in"},
                       {"image": "../charts/revenue.png", "duration": 2.5, "motion": "pan-left"}],
            "outro": {"text": "Thanks", "duration": 1.5},
        }
        (root / "video" / "video.json").write_text(json.dumps(vspec), encoding="utf-8")
        result = video.build(root / "video" / "video.json")
        expected = 2 + 2.5 + 2.5 + 1.5 - 3 * 0.5
        assert result["size"] == "1280x720", result
        assert abs(result["duration"] - expected) < 0.3, (result, expected)
        assert result["audio"], result
        assert "srt" in result
        print("video ok:", result)

        (root / "manifest.json").write_text(json.dumps({
            "title": "Smoke test run", "request": "Make one of everything", "created": "2026-10-10T12:00", "theme": "ocean",
            "deliverables": [{"title": "Deck", "agent": "slides-designer", "status": "done", "files": ["slides/deck.html"]},
                             {"title": "Photo", "agent": "image-maker", "status": "skipped", "note": "Needs OPENROUTER_API_KEY"}],
            "issues": ["Photo-realistic images need OPENROUTER_API_KEY"],
        }), encoding="utf-8")
        index = gallery.build(root)
        page = index.read_text(encoding="utf-8")
        assert "Smoke test run" in page and "deck.html" in page and "video.mp4" in page and "summary.docx" in page
        print("gallery ok:", index)
        assert set(THEMES) == {"midnight", "paper", "ocean", "forest", "ember"}
        print("\nALL OK", f"(kept at {root})" if keep else "")
        return 0
    finally:
        if not keep:
            shutil.rmtree(root, ignore_errors=True)


if __name__ == "__main__":
    sys.exit(main())
