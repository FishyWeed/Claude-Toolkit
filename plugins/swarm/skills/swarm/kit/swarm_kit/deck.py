"""One deck spec (JSON) -> a self-contained HTML deck, a matching .pptx and a PDF.

Spec:
{
  "title": "...", "subtitle": "...", "author": "...", "date": "...",
  "theme": "ocean|paper|midnight|forest|ember",
  "sources": ["[1] Title — Publisher, 2025. https://..."],
  "slides": [ {"layout": "...", ...}, ... ]
}

Layouts and their fields (every slide also takes "notes", "source", "kicker"):
  title        title, subtitle, image     (deck title/subtitle/author/date by default; image = full-bleed background)
  section      title, subtitle, image     (image = full-bleed background)
  bullets      title, bullets[]
  image        title, image, caption      (alias: chart)
  image-text   title, image, bullets[] or text
  two-column   title, left{heading, bullets[]}, right{heading, bullets[]}
  stats        title, stats[{value, label, source}]   (2-4 stats)
  quote        quote, attribution
  table        title, columns[], rows[[...]]
  closing      title, bullets[], image    (with an image it is laid out like image-text)
  sources      title                       (lists the deck's "sources")
Image paths are relative to the spec file.
"""

from __future__ import annotations

import base64
import html
import json
import mimetypes
from pathlib import Path

from . import inline
from .themes import Theme, get_theme, hex_to_rgb

W, H = 1920, 1080
MARGIN_X = 120
CONTENT_W = W - 2 * MARGIN_X

LAYOUTS = {"title", "section", "bullets", "image", "chart", "image-text", "two-column", "stats", "quote", "table", "closing", "sources"}


# ---------------------------------------------------------------- spec + checks


def load(spec_path: Path) -> dict:
    spec = json.loads(Path(spec_path).read_text(encoding="utf-8"))
    if not isinstance(spec.get("slides"), list) or not spec["slides"]:
        raise SystemExit("The deck spec needs a non-empty 'slides' list.")
    for i, slide in enumerate(spec["slides"], 1):
        layout = slide.get("layout", "bullets")
        if layout not in LAYOUTS:
            raise SystemExit(f"Slide {i}: unknown layout '{layout}'. Use one of: {', '.join(sorted(LAYOUTS))}")
        slide["layout"] = "image" if layout == "chart" else layout
        if slide["layout"] == "closing" and slide.get("image"):
            slide["layout"] = "image-text"
    return spec


def check(spec: dict, base: Path) -> list[str]:
    """Warnings a designer should fix: they do not stop the build."""
    warnings = []
    for i, s in enumerate(spec["slides"], 1):
        title = inline.plain(s.get("title", ""))
        if len(title.split()) > 14:
            warnings.append(f"Slide {i}: title has {len(title.split())} words; keep it under ~12.")
        bullets = s.get("bullets") or []
        if len(bullets) > 6:
            warnings.append(f"Slide {i}: {len(bullets)} bullets; keep it to 6 or fewer (split the slide).")
        for b in bullets:
            if len(inline.plain(b).split()) > 22:
                warnings.append(f"Slide {i}: a bullet runs {len(inline.plain(b).split())} words; tighten it or move detail to notes.")
                break
        if s["layout"] in {"image", "image-text"}:
            img = s.get("image")
            if not img:
                warnings.append(f"Slide {i}: layout '{s['layout']}' has no image.")
            elif not (base / img).exists():
                warnings.append(f"Slide {i}: image not found: {img}")
        if s["layout"] == "stats" and not (2 <= len(s.get("stats") or []) <= 4):
            warnings.append(f"Slide {i}: stats layout wants 2-4 stats.")
    if any(s["layout"] == "sources" for s in spec["slides"]) and not spec.get("sources"):
        warnings.append("A 'sources' slide is present but the deck has no 'sources' list.")
    return warnings


def _bullet_px(bullets: list[str]) -> int:
    chars = sum(len(inline.plain(b)) for b in bullets)
    if len(bullets) <= 3 and chars <= 120:
        return 48
    if len(bullets) <= 4 and chars <= 200:
        return 40
    if chars <= 320:
        return 36
    if chars <= 440:
        return 32
    return 28


def _image_size(path: Path) -> tuple[int, int]:
    try:
        from PIL import Image

        with Image.open(path) as im:
            return im.size
    except Exception:
        return (16, 9)


def _fit(box_w: int, box_h: int, img_w: int, img_h: int) -> tuple[int, int, int, int]:
    """Contain: returns (dx, dy, w, h) inside the box, centred."""
    scale = min(box_w / img_w, box_h / img_h)
    w, h = int(img_w * scale), int(img_h * scale)
    return (box_w - w) // 2, (box_h - h) // 2, w, h


# ---------------------------------------------------------------- HTML


def _data_uri(path: Path) -> str:
    mime = mimetypes.guess_type(path.name)[0] or "image/png"
    return f"data:{mime};base64,{base64.b64encode(path.read_bytes()).decode()}"


def _ul(bullets: list[str], px: int) -> str:
    items = "".join(f"<li>{inline.to_html(b)}</li>" for b in bullets)
    return f'<ul class="bullets" style="font-size:{px}px">{items}</ul>'


def _header(s: dict) -> str:
    kicker = f'<div class="kicker">{inline.to_html(s["kicker"])}</div>' if s.get("kicker") else ""
    return f'<header>{kicker}<h2>{inline.to_html(s.get("title", ""))}</h2></header>'


def _slide_html(s: dict, spec: dict, base: Path, n: int, total: int) -> str:
    layout = s["layout"]
    body = ""
    cls = layout
    style = ""
    bg = base / s["image"] if layout in {"title", "section"} and s.get("image") else None
    if bg is not None and bg.is_file():
        cls += " has-bg"
        style = f' style="background:url({_data_uri(bg)}) center/cover no-repeat"'
    if layout == "title":
        kicker = f'<div class="kicker">{inline.to_html(s["kicker"])}</div>' if s.get("kicker") else ""
        meta = " · ".join(x for x in [spec.get("author", ""), spec.get("date", "")] if x)
        body = (
            f'<div class="title-block">{kicker}<h1>{inline.to_html(s.get("title") or spec.get("title", ""))}</h1>'
            f'<div class="bar"></div><p class="subtitle">{inline.to_html(s.get("subtitle") or spec.get("subtitle", ""))}</p></div>'
            f'<div class="meta">{html.escape(meta)}</div>'
        )
    elif layout == "section":
        kicker = f'<div class="kicker">{inline.to_html(s["kicker"])}</div>' if s.get("kicker") else ""
        sub = f'<p class="subtitle">{inline.to_html(s["subtitle"])}</p>' if s.get("subtitle") else ""
        body = f'<div class="title-block">{kicker}<h1>{inline.to_html(s.get("title", ""))}</h1>{sub}</div>'
    elif layout in {"bullets", "closing"}:
        body = _header(s) + _ul(s.get("bullets") or [], _bullet_px(s.get("bullets") or []))
    elif layout == "image":
        img = base / s.get("image", "")
        figure = f'<img src="{_data_uri(img)}" alt="">' if img.is_file() else '<div class="missing">image missing</div>'
        cap = f'<figcaption>{inline.to_html(s["caption"])}</figcaption>' if s.get("caption") else ""
        body = _header(s) + f'<figure>{figure}{cap}</figure>'
    elif layout == "image-text":
        img = base / s.get("image", "")
        figure = f'<img src="{_data_uri(img)}" alt="">' if img.is_file() else '<div class="missing">image missing</div>'
        text = _ul(s["bullets"], min(_bullet_px(s["bullets"]), 34)) if s.get("bullets") else f'<p class="lead">{inline.to_html(s.get("text", ""))}</p>'
        body = f'<div class="split"><div class="split-img">{figure}</div><div class="split-text">{_header(s)}{text}</div></div>'
    elif layout == "two-column":
        cols = []
        for side in ("left", "right"):
            c = s.get(side) or {}
            items = c.get("bullets") or []
            cols.append(f'<div class="card"><h3>{inline.to_html(c.get("heading", ""))}</h3>{_ul(items, min(_bullet_px(items), 32))}</div>')
        body = _header(s) + f'<div class="cols">{"".join(cols)}</div>'
    elif layout == "stats":
        cards = []
        for st in s.get("stats") or []:
            src = f'<div class="stat-src">{inline.to_html(st["source"])}</div>' if st.get("source") else ""
            cards.append(f'<div class="stat"><div class="stat-value">{inline.to_html(str(st.get("value", "")))}</div>'
                         f'<div class="stat-label">{inline.to_html(st.get("label", ""))}</div>{src}</div>')
        body = _header(s) + f'<div class="stats">{"".join(cards)}</div>'
    elif layout == "quote":
        attr = f'<div class="attribution">— {inline.to_html(s["attribution"])}</div>' if s.get("attribution") else ""
        body = f'<div class="quote-block"><div class="qmark">“</div><blockquote>{inline.to_html(s.get("quote", ""))}</blockquote>{attr}</div>'
    elif layout == "table":
        cols = s.get("columns") or []
        rows = s.get("rows") or []
        px = 30 if len(rows) <= 5 else 26 if len(rows) <= 8 else 22
        head = "".join(f"<th>{inline.to_html(str(c))}</th>" for c in cols)
        trs = "".join("<tr>" + "".join(f"<td>{inline.to_html(str(c))}</td>" for c in r) + "</tr>" for r in rows)
        body = _header(s) + f'<table style="font-size:{px}px"><thead><tr>{head}</tr></thead><tbody>{trs}</tbody></table>'
    elif layout == "sources":
        s = {**s, "title": s.get("title") or "Sources"}
        items = "".join(f"<li>{inline.to_html(x)}</li>" for x in spec.get("sources") or [])
        body = _header(s) + f'<ol class="sources">{items}</ol>'
    footer = ""
    if layout not in {"title", "section"}:
        src = f'<span class="src">{inline.to_html(s["source"])}</span>' if s.get("source") else "<span></span>"
        footer = f'<footer>{src}<span class="num">{n} / {total}</span></footer>'
    notes = f'<aside class="notes">{html.escape(s.get("notes", ""))}</aside>' if s.get("notes") else ""
    return f'<section class="slide layout-{cls}"{style} data-n="{n}">{body}{footer}{notes}</section>'


def build_html(spec: dict, base: Path, theme: Theme) -> str:
    slides = spec["slides"]
    sections = "\n".join(_slide_html(s, spec, base, i, len(slides)) for i, s in enumerate(slides, 1))
    t = theme
    shade = ",".join(str(c) for c in hex_to_rgb(t.bg if t.dark else t.text))  # the theme's darkest colour
    css = f"""
:root{{--bg:{t.bg};--surface:{t.surface};--text:{t.text};--muted:{t.muted};--accent:{t.accent};--accent2:{t.accent2};--on-accent:{t.on_accent}}}
*{{box-sizing:border-box}}
html,body{{margin:0;height:100%;background:#0b0b0c;overflow:hidden}}
.slide{{position:absolute;left:50%;top:50%;width:{W}px;height:{H}px;margin:-{H//2}px 0 0 -{W//2}px;display:none;flex-direction:column;
  padding:96px {MARGIN_X}px 120px;background:var(--bg);color:var(--text);font-family:{t.body_stack};overflow:hidden}}
.slide.active{{display:flex}}
h1,h2,h3,blockquote,.stat-value{{font-family:{t.heading_stack};font-weight:700;margin:0;letter-spacing:-0.01em}}
h2{{font-size:64px;line-height:1.1;max-width:1600px}}
h3{{font-size:36px;margin-bottom:24px;color:var(--accent)}}
.kicker{{font-size:24px;letter-spacing:.14em;text-transform:uppercase;color:var(--accent);font-weight:700;margin-bottom:18px}}
header{{margin-bottom:56px}}
header::after{{content:"";display:block;width:120px;height:8px;background:var(--accent);margin-top:28px;border-radius:4px}}
a{{color:var(--accent)}}
code{{font-family:Consolas,monospace;background:var(--surface);padding:0 .25em;border-radius:6px}}
.bullets{{list-style:none;margin:0;padding:0;line-height:1.35}}
.bullets li{{position:relative;padding-left:1.4em;margin-bottom:.7em}}
.bullets li::before{{content:"";position:absolute;left:0;top:.42em;width:.5em;height:.5em;background:var(--accent);border-radius:3px}}
.layout-title .title-block,.layout-section .title-block{{margin:auto 0}}
.layout-title h1{{font-size:104px;line-height:1.05;max-width:1500px}}
.layout-title .bar{{width:160px;height:10px;background:var(--accent);border-radius:5px;margin:48px 0 40px}}
.layout-title .subtitle{{font-size:42px;color:var(--muted);margin:0;max-width:1400px;line-height:1.3}}
.layout-title .meta{{font-size:26px;color:var(--muted)}}
.layout-section{{background:var(--accent);color:var(--on-accent)}}
.layout-section .kicker{{color:var(--on-accent);opacity:.8}}
.layout-section h1{{font-size:96px;line-height:1.05;max-width:1500px}}
.layout-section .subtitle{{font-size:40px;margin-top:32px;opacity:.85}}
figure{{margin:0;flex:1;display:flex;flex-direction:column;align-items:center;justify-content:center;min-height:0}}
figure img{{max-width:100%;max-height:100%;min-height:0;object-fit:contain;border-radius:12px}}
figcaption{{font-size:24px;color:var(--muted);margin-top:20px}}
.missing{{padding:80px;border:4px dashed var(--muted);color:var(--muted);font-size:32px;border-radius:16px}}
.split{{display:grid;grid-template-columns:1fr 1fr;gap:80px;flex:1;min-height:0}}
.split-img{{display:flex;align-items:center;justify-content:center;min-height:0}}
.split-img img{{max-width:100%;max-height:100%;object-fit:contain;border-radius:16px}}
.split-text{{display:flex;flex-direction:column;justify-content:center}}
.lead{{font-size:34px;line-height:1.45;margin:0}}
.cols{{display:grid;grid-template-columns:1fr 1fr;gap:48px;flex:1;min-height:0}}
.card{{background:var(--surface);border-radius:24px;padding:48px}}
.stats{{display:flex;gap:40px;flex:1;align-items:center}}
.stat{{flex:1;background:var(--surface);border-radius:24px;padding:56px 48px;border-top:10px solid var(--accent)}}
.stat:nth-child(even){{border-top-color:var(--accent2)}}
.stat-value{{font-size:112px;line-height:1;color:var(--accent)}}
.stat:nth-child(even) .stat-value{{color:var(--accent2)}}
.stat-label{{font-size:32px;margin-top:24px;line-height:1.3}}
.stat-src{{font-size:20px;color:var(--muted);margin-top:20px}}
.quote-block{{margin:auto 0;max-width:1500px}}
.qmark{{font-family:Georgia,serif;font-size:240px;line-height:.6;color:var(--accent);height:110px}}
blockquote{{font-size:64px;line-height:1.25;font-style:italic;font-weight:400}}
.attribution{{font-size:32px;color:var(--muted);margin-top:48px}}
table{{border-collapse:collapse;width:100%}}
th{{background:var(--accent);color:var(--on-accent);text-align:left;padding:.6em .8em;font-weight:700}}
td{{padding:.55em .8em;border-bottom:2px solid var(--surface)}}
tbody tr:nth-child(even) td{{background:var(--surface)}}
.sources{{columns:2;column-gap:80px;font-size:26px;line-height:1.4;color:var(--muted);margin:0;padding-left:1.4em}}
.sources li{{margin-bottom:.6em;break-inside:avoid}}
footer{{position:absolute;left:{MARGIN_X}px;right:{MARGIN_X}px;bottom:44px;display:flex;justify-content:space-between;align-items:flex-end;gap:40px;font-size:20px;color:var(--muted)}}
.layout-section footer{{color:var(--on-accent)}}
.has-bg{{color:#fff}}
.has-bg::before{{content:"";position:absolute;inset:0;background:linear-gradient(90deg,rgba({shade},.9) 0%,rgba({shade},.72) 38%,rgba({shade},0) 70%)}}
.has-bg>*{{position:relative;z-index:1}}
.has-bg h1{{color:#fff;max-width:1050px}}
.has-bg .kicker{{color:#fff;opacity:.85}}
.has-bg .bar{{background:var(--accent2)}}
.has-bg .subtitle,.has-bg .meta{{color:rgba(255,255,255,.86)!important;opacity:1;max-width:900px}}
.notes{{display:none}}
#hud{{position:fixed;right:16px;bottom:12px;font:13px system-ui,sans-serif;color:#9aa;opacity:.7;user-select:none}}
@page{{size:{W}px {H}px;margin:0}}
@media print{{
  html,body{{background:none;overflow:visible;height:auto}}
  .slide{{display:flex!important;position:relative;left:0;top:0;margin:0;transform:none!important;page-break-after:always;break-after:page}}
  #hud{{display:none}}
}}
"""
    js = """
const slides=[...document.querySelectorAll('.slide')];let i=0;
const hud=document.getElementById('hud');
function fit(){const s=Math.min(innerWidth/%d,innerHeight/%d);slides.forEach(el=>el.style.transform='scale('+s+')')}
function show(n){i=Math.max(0,Math.min(slides.length-1,n));slides.forEach((el,k)=>el.classList.toggle('active',k===i));hud.textContent=(i+1)+' / '+slides.length;history.replaceState(null,'','#'+(i+1))}
addEventListener('keydown',e=>{if(['ArrowRight','PageDown',' ','Enter'].includes(e.key)){show(i+1);e.preventDefault()}
 else if(['ArrowLeft','PageUp','Backspace'].includes(e.key)){show(i-1);e.preventDefault()}
 else if(e.key==='Home')show(0);else if(e.key==='End')show(slides.length-1);else if(e.key==='f')document.documentElement.requestFullscreen?.()});
addEventListener('click',e=>{if(e.target.closest('a'))return;show(e.clientX>innerWidth/3?i+1:i-1)});
if(window.top!==window)hud.style.display='none';
addEventListener('resize',fit);fit();show((parseInt(location.hash.slice(1))||1)-1);
""" % (W, H)
    title = html.escape(spec.get("title", "Deck"))
    return (f"<!doctype html><html lang='en'><head><meta charset='utf-8'><meta name='viewport' content='width=device-width,initial-scale=1'>"
            f"<title>{title}</title><style>{css}</style></head><body>{sections}<div id='hud'></div><script>{js}</script></body></html>")


# ---------------------------------------------------------------- PPTX


def build_pptx(spec: dict, base: Path, theme: Theme, out: Path) -> None:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.oxml.ns import qn
    from pptx.util import Emu, Pt
    from lxml import etree

    EMU_PX = 6350  # 12192000 EMU / 1920 px

    def px(v: float) -> Emu:
        return Emu(int(v * EMU_PX))

    def rgb(hexv: str) -> RGBColor:
        return RGBColor(*hex_to_rgb(hexv))

    def fpt(px_size: float) -> Pt:
        return Pt(px_size * 0.5)  # 1920 px = 960 pt

    prs = Presentation()
    prs.slide_width, prs.slide_height = px(W), px(H)
    blank = prs.slide_layouts[6]
    t = theme

    def rect(slide, x, y, w, h, color, rounded=False):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE, px(x), px(y), px(w), px(h))
        shape.fill.solid()
        shape.fill.fore_color.rgb = rgb(color)
        shape.line.fill.background()
        shape.shadow.inherit = False
        if rounded:
            shape.adjustments[0] = min(0.5, 24 / max(1, min(w, h)))
        return shape

    def add_runs(paragraph, text, size, color, font, bold=False, italic=False):
        for r in inline.parse(text, bold=bold, italic=italic):
            run = paragraph.add_run()
            run.text = r.text
            f = run.font
            f.size = fpt(size)
            f.bold = r.bold
            f.italic = r.italic
            f.name = "Consolas" if r.code else font
            f.color.rgb = rgb(t.accent if r.link else color)
            if r.link:
                run.hyperlink.address = r.link

    def textbox(slide, x, y, w, h, text, size, color=None, font=None, bold=False, italic=False,
                align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, line=1.15):
        box = slide.shapes.add_textbox(px(x), px(y), px(w), px(h))
        tf = box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = anchor
        lines = text if isinstance(text, list) else [text]
        for k, ln in enumerate(lines):
            p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
            p.alignment = align
            p.line_spacing = line
            add_runs(p, ln, size, color or t.text, font or t.body_font, bold, italic)
        return box

    def bullets(slide, x, y, w, h, items, size):
        box = slide.shapes.add_textbox(px(x), px(y), px(w), px(h))
        tf = box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        for k, item in enumerate(items):
            p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
            p.line_spacing = 1.2
            p.space_after = fpt(size * 0.6)
            pPr = p._p.get_or_add_pPr()
            pPr.set("marL", str(int(size * 1.4 * EMU_PX)))
            pPr.set("indent", str(-int(size * 1.4 * EMU_PX)))
            for tag in ("a:buNone", "a:buChar", "a:buClr", "a:buFont", "a:buSzPct"):
                for old in pPr.findall(qn(tag)):
                    pPr.remove(old)
            clr = etree.SubElement(pPr, qn("a:buClr"))
            etree.SubElement(clr, qn("a:srgbClr")).set("val", t.accent.lstrip("#"))
            etree.SubElement(pPr, qn("a:buSzPct")).set("val", "100000")
            etree.SubElement(pPr, qn("a:buFont")).set("typeface", "Arial")
            etree.SubElement(pPr, qn("a:buChar")).set("char", "■")
            add_runs(p, item, size, t.text, t.body_font)
        return box

    def header(slide, s) -> int:
        y = 96
        if s.get("kicker"):
            textbox(slide, MARGIN_X, y, CONTENT_W, 34, inline.plain(s["kicker"]).upper(), 24, t.accent, bold=True)
            y += 52
        title = s.get("title", "")
        lines = 2 if len(inline.plain(title)) > 44 else 1
        textbox(slide, MARGIN_X, y, 1600, 76 * lines, title, 64, t.text, t.heading_font, bold=True, line=1.05)
        y += 76 * lines + 28
        rect(slide, MARGIN_X, y, 120, 8, t.accent)
        return y + 56

    def picture(slide, path: Path, x, y, w, h):
        if not path.is_file():
            textbox(slide, x, y, w, h, "image missing", 32, t.muted, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
            return
        iw, ih = _image_size(path)
        dx, dy, fw, fh = _fit(w, h, iw, ih)
        slide.shapes.add_picture(str(path), px(x + dx), px(y + dy), px(fw), px(fh))

    def background(slide, path: Path) -> bool:
        """Full-bleed picture, cover-cropped, under a left-to-right dark gradient for legible white text."""
        if not path.is_file():
            return False
        iw, ih = _image_size(path)
        pic = slide.shapes.add_picture(str(path), 0, 0, px(W), px(H))
        if iw / ih > W / H:
            excess = 1 - (W / H) / (iw / ih)
            pic.crop_left = pic.crop_right = excess / 2
        else:
            excess = 1 - (iw / ih) / (W / H)
            pic.crop_top = pic.crop_bottom = excess / 2
        shade = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, px(W), px(H))
        shade.line.fill.background()
        shade.shadow.inherit = False
        shade.fill.gradient()
        shade.fill.gradient_angle = 0
        tint = rgb(t.bg if t.dark else t.text)
        for stop, pos, alpha in zip(shade.fill.gradient_stops, (0.0, 0.70), (90000, 0)):
            stop.position = pos
            stop.color.rgb = tint
            clr = stop._element.find(qn("a:srgbClr"))
            etree.SubElement(clr, qn("a:alpha")).set("val", str(alpha))
        return True

    total = len(spec["slides"])
    bottom = H - 120
    for n, s in enumerate(spec["slides"], 1):
        layout = s["layout"]
        slide = prs.slides.add_slide(blank)
        fill = slide.background.fill
        fill.solid()
        fill.fore_color.rgb = rgb(t.accent if layout == "section" else t.bg)

        has_bg = layout in {"title", "section"} and bool(s.get("image")) and background(slide, base / s["image"])
        if layout == "title":
            ink, soft, kick = ("#FFFFFF", "#E6E6E6", "#FFFFFF") if has_bg else (t.text, t.muted, t.accent)
            width = 1050 if has_bg else 1500
            y = 300
            if s.get("kicker"):
                textbox(slide, MARGIN_X, y - 60, CONTENT_W, 34, inline.plain(s["kicker"]).upper(), 24, kick, bold=True)
            title = s.get("title") or spec.get("title", "")
            lines = max(1, min(3, -(-len(inline.plain(title)) // (18 if has_bg else 28))))
            textbox(slide, MARGIN_X, y, width, 112 * lines, title, 104, ink, t.heading_font, bold=True, line=1.02)
            y += 112 * lines + 48
            rect(slide, MARGIN_X, y, 160, 10, t.accent2 if has_bg else t.accent)
            sub = s.get("subtitle") or spec.get("subtitle", "")
            if sub:
                textbox(slide, MARGIN_X, y + 50, 900 if has_bg else 1400, 120, sub, 42, soft)
            meta = " · ".join(x for x in [spec.get("author", ""), spec.get("date", "")] if x)
            if meta:
                textbox(slide, MARGIN_X, H - 140, CONTENT_W, 40, meta, 26, soft)
        elif layout == "section":
            ink = "#FFFFFF" if has_bg else t.on_accent
            if s.get("kicker"):
                textbox(slide, MARGIN_X, 330, CONTENT_W, 34, inline.plain(s["kicker"]).upper(), 24, ink, bold=True)
            textbox(slide, MARGIN_X, 390, 1050 if has_bg else 1500, 220, s.get("title", ""), 96, ink, t.heading_font, bold=True, line=1.02)
            if s.get("subtitle"):
                textbox(slide, MARGIN_X, 640, 1050 if has_bg else 1500, 120, s["subtitle"], 40, ink)
        elif layout in {"bullets", "closing"}:
            y = header(slide, s)
            items = s.get("bullets") or []
            bullets(slide, MARGIN_X, y, CONTENT_W, bottom - y, items, _bullet_px(items))
        elif layout == "image":
            y = header(slide, s)
            cap_h = 50 if s.get("caption") else 0
            picture(slide, base / s.get("image", ""), MARGIN_X, y, CONTENT_W, bottom - y - cap_h)
            if s.get("caption"):
                textbox(slide, MARGIN_X, bottom - 36, CONTENT_W, 36, s["caption"], 24, t.muted, align=PP_ALIGN.CENTER)
        elif layout == "image-text":
            half = (CONTENT_W - 80) // 2
            picture(slide, base / s.get("image", ""), MARGIN_X, 96, half, bottom - 96)
            sub = dict(s)
            box_x = MARGIN_X + half + 80
            y = 200
            if s.get("kicker"):
                textbox(slide, box_x, y, half, 34, inline.plain(s["kicker"]).upper(), 24, t.accent, bold=True)
                y += 52
            textbox(slide, box_x, y, half, 160, sub.get("title", ""), 56, t.text, t.heading_font, bold=True, line=1.05)
            y += 190
            rect(slide, box_x, y, 120, 8, t.accent)
            y += 56
            if s.get("bullets"):
                bullets(slide, box_x, y, half, bottom - y, s["bullets"], min(_bullet_px(s["bullets"]), 34))
            else:
                textbox(slide, box_x, y, half, bottom - y, s.get("text", ""), 34, t.text, line=1.4)
        elif layout == "two-column":
            y = header(slide, s)
            col_w = (CONTENT_W - 48) // 2
            for k, side in enumerate(("left", "right")):
                c = s.get(side) or {}
                x = MARGIN_X + k * (col_w + 48)
                rect(slide, x, y, col_w, bottom - y, t.surface, rounded=True)
                textbox(slide, x + 48, y + 44, col_w - 96, 50, c.get("heading", ""), 36, t.accent, t.heading_font, bold=True)
                items = c.get("bullets") or []
                bullets(slide, x + 48, y + 124, col_w - 96, bottom - y - 160, items, min(_bullet_px(items), 32))
        elif layout == "stats":
            y = header(slide, s)
            stats = s.get("stats") or []
            k = max(1, len(stats))
            gap = 40
            card_w = (CONTENT_W - gap * (k - 1)) // k
            card_h = min(440, bottom - y)
            cy = y + (bottom - y - card_h) // 2
            for j, st in enumerate(stats):
                x = MARGIN_X + j * (card_w + gap)
                color = t.accent if j % 2 == 0 else t.accent2
                rect(slide, x, cy, card_w, card_h, t.surface, rounded=True)
                rect(slide, x, cy, card_w, 10, color)
                textbox(slide, x + 48, cy + 64, card_w - 96, 130, str(st.get("value", "")), 112 if k <= 3 else 88, color, t.heading_font, bold=True)
                textbox(slide, x + 48, cy + 220, card_w - 96, 120, st.get("label", ""), 32, t.text)
                if st.get("source"):
                    textbox(slide, x + 48, cy + card_h - 70, card_w - 96, 40, st["source"], 20, t.muted)
        elif layout == "quote":
            textbox(slide, MARGIN_X, 170, 200, 200, "“", 240, t.accent, "Georgia", bold=True)
            textbox(slide, MARGIN_X, 340, 1500, 420, s.get("quote", ""), 64, t.text, t.heading_font, italic=True, line=1.2)
            if s.get("attribution"):
                textbox(slide, MARGIN_X, 800, 1500, 50, "— " + s["attribution"], 32, t.muted)
        elif layout == "table":
            y = header(slide, s)
            cols = s.get("columns") or []
            rows = s.get("rows") or []
            size = 30 if len(rows) <= 5 else 26 if len(rows) <= 8 else 22
            row_h = int(size * 2.1)
            if cols:
                shape = slide.shapes.add_table(len(rows) + 1, len(cols), px(MARGIN_X), px(y), px(CONTENT_W), px(row_h * (len(rows) + 1)))
                table = shape.table
                tbl_pr = shape._element.graphic.graphicData.tbl.tblPr
                for attr in ("bandRow", "firstRow"):
                    tbl_pr.set(attr, "0")
                for r_i in range(len(rows) + 1):
                    for c_i in range(len(cols)):
                        cell = table.cell(r_i, c_i)
                        value = cols[c_i] if r_i == 0 else (rows[r_i - 1][c_i] if c_i < len(rows[r_i - 1]) else "")
                        cell.fill.solid()
                        cell.fill.fore_color.rgb = rgb(t.accent if r_i == 0 else (t.surface if r_i % 2 == 0 else t.bg))
                        cell.margin_left = cell.margin_right = px(size * 0.8)
                        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                        tf = cell.text_frame
                        tf.paragraphs[0].text = ""
                        add_runs(tf.paragraphs[0], str(value), size, t.on_accent if r_i == 0 else t.text, t.body_font, bold=(r_i == 0))
        elif layout == "sources":
            y = header(slide, {**s, "title": s.get("title") or "Sources"})
            items = spec.get("sources") or []
            half = (len(items) + 1) // 2
            col_w = (CONTENT_W - 80) // 2
            for k, chunk in enumerate((items[:half], items[half:])):
                if chunk:
                    textbox(slide, MARGIN_X + k * (col_w + 80), y, col_w, bottom - y, list(chunk), 26, t.muted, line=1.3)

        if layout not in {"title", "section"}:
            if s.get("source"):
                textbox(slide, MARGIN_X, H - 72, 1500, 32, s["source"], 20, t.muted)
            textbox(slide, W - MARGIN_X - 200, H - 72, 200, 32, f"{n} / {total}", 20, t.muted, align=PP_ALIGN.RIGHT)
        if s.get("notes"):
            slide.notes_slide.notes_text_frame.text = s["notes"]

    prs.core_properties.title = spec.get("title", "")
    prs.core_properties.author = spec.get("author", "")
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out))


# ---------------------------------------------------------------- entry point


def build(spec_path: Path, out_dir: Path | None = None, name: str = "deck", pdf: bool = True) -> dict:
    spec_path = Path(spec_path).resolve()
    base = spec_path.parent
    out_dir = Path(out_dir).resolve() if out_dir else base
    spec = load(spec_path)
    theme = get_theme(spec.get("theme"))
    warnings = check(spec, base)

    html_path = out_dir / f"{name}.html"
    out_dir.mkdir(parents=True, exist_ok=True)
    html_path.write_text(build_html(spec, base, theme), encoding="utf-8")
    pptx_path = out_dir / f"{name}.pptx"
    build_pptx(spec, base, theme, pptx_path)
    result = {"html": str(html_path), "pptx": str(pptx_path), "slides": len(spec["slides"]), "theme": theme.name, "warnings": warnings}
    if pdf:
        from .browser import html_to_pdf

        try:
            result["pdf"] = str(html_to_pdf(html_path, out_dir / f"{name}.pdf"))
        except Exception as exc:  # a missing browser should not lose the deck
            result["warnings"].append(f"PDF not made: {exc}")
    return result


def preview(deck_html: Path, out_png: Path | None = None, cols: int = 3, tile_w: int = 640) -> dict:
    """Renders every slide of a built deck into one contact sheet PNG, for a visual check."""
    from PIL import Image, ImageDraw

    from .browser import screenshot

    deck_html = Path(deck_html).resolve()
    count = deck_html.read_text(encoding="utf-8").count('<section class="slide ')
    out_png = Path(out_png).resolve() if out_png else deck_html.with_name(deck_html.stem + "-preview.png")
    tile_h = tile_w * H // W
    gap = 16
    rows = (count + cols - 1) // cols
    sheet = Image.new("RGB", (cols * tile_w + (cols + 1) * gap, rows * tile_h + (rows + 1) * gap), (40, 40, 44))
    draw = ImageDraw.Draw(sheet)
    import tempfile

    with tempfile.TemporaryDirectory(prefix="swarm-preview-") as tmp:
        for n in range(1, count + 1):
            shot = screenshot(deck_html, Path(tmp) / f"s{n}.png", W // 2, H // 2, fragment=f"#{n}")
            with Image.open(shot) as im:
                tile = im.convert("RGB").resize((tile_w, tile_h), Image.LANCZOS)
            x = gap + ((n - 1) % cols) * (tile_w + gap)
            y = gap + ((n - 1) // cols) * (tile_h + gap)
            sheet.paste(tile, (x, y))
            draw.rectangle([x, y, x + 34, y + 22], fill=(0, 0, 0))
            draw.text((x + 6, y + 5), str(n), fill=(255, 255, 255))
    sheet.save(out_png)
    return {"preview": str(out_png), "slides": count}
